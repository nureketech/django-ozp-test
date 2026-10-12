from datetime import timedelta
from io import BytesIO, StringIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.quiz.models import ExamSession, Subject

from .models import Profile, StudyGroup
from .student_import import iin_is_valid, normalize_iin, split_name


def informatics():
    """Деректер миграциясы жасаған «Информатика» пәні."""
    return Subject.objects.get(code="informatics")


class ProfileTests(TestCase):
    def test_profile_created_for_new_user(self):
        user = User.objects.create_user(username="student", password="pass12345")
        self.assertTrue(Profile.objects.filter(user=user).exists())
        self.assertIsNone(user.profile.group)

    def test_profile_group(self):
        group = StudyGroup.objects.create(name="ИНФ-21", subject=informatics(), language="kk")
        user = User.objects.create_user(username="student", password="pass12345")
        user.profile.group = group
        user.profile.save()
        self.assertEqual(group.profiles.get().user, user)
        self.assertEqual(str(group), "ИНФ-21")


# ---------- 3-кезең: тіркелу, кіру, шығу ----------


class RegistrationRemovedTests(TestCase):
    """Сайтта тіркелу жоқ: студенттерді әкімші admin-де импорттайды."""

    def test_register_url_is_gone(self):
        self.assertEqual(self.client.get("/accounts/register/").status_code, 404)

    def test_no_register_links(self):
        for url in [reverse("quiz:home"), reverse("accounts:login")]:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertNotContains(response, "/accounts/register/")
                self.assertNotContains(response, "Тіркелу")


class LoginLogoutTests(TestCase):
    def setUp(self):
        self.student = User.objects.create_user(username="student", password="pass12345")
        self.teacher = User.objects.create_user(
            username="teacher", password="pass12345", is_staff=True
        )
        self.teacher.profile.subjects.add(informatics())

    def login(self, username, next_url=None):
        url = reverse("accounts:login")
        if next_url:
            url += f"?next={next_url}"
        return self.client.post(url, {"username": username, "password": "pass12345"})

    def test_student_goes_to_dashboard(self):
        response = self.login("student")
        self.assertRedirects(response, reverse("accounts:after_login"), fetch_redirect_response=False)
        response = self.client.get(response["Location"])
        self.assertRedirects(response, reverse("quiz:dashboard"))

    def test_teacher_goes_to_questions(self):
        response = self.client.get(self.login("teacher")["Location"])
        self.assertRedirects(response, reverse("quiz:teacher_questions"))

    def test_next_parameter_is_respected(self):
        response = self.login("student", next_url=reverse("quiz:dashboard"))
        self.assertRedirects(response, reverse("quiz:dashboard"))

    def test_wrong_password(self):
        response = self.client.post(
            reverse("accounts:login"), {"username": "student", "password": "wrong"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ИИН (логин) немесе код (құпия сөз) қате.")
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_logout_requires_post_and_goes_home(self):
        self.client.force_login(self.student)
        response = self.client.post(reverse("accounts:logout"))
        self.assertRedirects(response, reverse("quiz:home"))
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_pages_open_in_russian(self):
        self.client.cookies["django_language"] = "ru"
        self.assertContains(self.client.get(reverse("accounts:login")), 'lang="ru"')


class NavigationTests(TestCase):
    def test_guest_sees_login(self):
        response = self.client.get(reverse("quiz:home"))
        self.assertContains(response, reverse("accounts:login"))
        self.assertNotContains(response, reverse("accounts:logout"))

    def test_student_sees_dashboard_link(self):
        student = User.objects.create_user(username="student", password="pass12345")
        self.client.force_login(student)
        response = self.client.get(reverse("quiz:home"))
        self.assertContains(response, reverse("quiz:dashboard"))
        self.assertContains(response, reverse("accounts:logout"))
        self.assertNotContains(response, reverse("quiz:teacher_questions"))
        self.assertNotContains(response, reverse("admin:index"))

    def test_teacher_sees_teacher_links(self):
        teacher = User.objects.create_user(
            username="teacher", password="pass12345", is_staff=True
        )
        self.client.force_login(teacher)
        response = self.client.get(reverse("quiz:home"))
        for url_name in [
            "quiz:teacher_questions",
            "quiz:teacher_contexts",
            "quiz:teacher_bank",
            "admin:index",
            "accounts:logout",
        ]:
            with self.subTest(url_name=url_name):
                self.assertContains(response, reverse(url_name))
        self.assertNotContains(response, reverse("quiz:dashboard"))


# ---------- Оқытушыны қосу: тек admin арқылы ----------


class AdminAddTeacherTests(TestCase):
    """Әкімші оқытушыны admin-де бір бетте қосады: логин, аты-жөні, is_staff, пәндері."""

    def setUp(self):
        self.admin = User.objects.create_superuser(username="admin", password="pass12345")
        self.client.force_login(self.admin)
        self.subject = informatics()

    def add_user(self, **overrides):
        data = {
            "username": "teacher",
            "usable_password": "true",
            "password1": "Qazaq-Test-2026",
            "password2": "Qazaq-Test-2026",
            "first_name": "Ерлан",
            "last_name": "Серікұлы",
            "is_staff": "on",
            "profile-TOTAL_FORMS": "1",
            "profile-INITIAL_FORMS": "0",
            "profile-MIN_NUM_FORMS": "0",
            "profile-MAX_NUM_FORMS": "1",
            "profile-0-group": "",
            "profile-0-subjects": [self.subject.pk],
            "_save": "Save",
        }
        data.update(overrides)
        return self.client.post(reverse("admin:auth_user_add"), data)

    def test_add_page_has_teacher_fields(self):
        response = self.client.get(reverse("admin:auth_user_add"))
        self.assertContains(response, 'name="first_name"')
        self.assertContains(response, 'name="is_staff"')
        self.assertContains(response, 'name="profile-0-subjects"')

    def test_teacher_with_subject_in_one_step(self):
        response = self.add_user()
        self.assertEqual(response.status_code, 302)
        teacher = User.objects.get(username="teacher")
        self.assertTrue(teacher.is_staff)
        self.assertFalse(teacher.is_superuser)
        self.assertEqual(teacher.get_full_name(), "Ерлан Серікұлы")
        self.assertEqual(list(teacher.profile.subjects.all()), [self.subject])
        self.assertEqual(Profile.objects.filter(user=teacher).count(), 1)

        # Жаңа оқытушы кіріп, өз пәнінің беттерін көреді
        self.client.force_login(teacher)
        response = self.client.get(reverse("quiz:teacher_questions"))
        self.assertEqual(response.status_code, 200)

    def test_user_without_profile_changes(self):
        response = self.add_user(**{"profile-0-subjects": [], "is_staff": ""})
        self.assertEqual(response.status_code, 302)
        user = User.objects.get(username="teacher")
        self.assertFalse(user.is_staff)
        self.assertEqual(Profile.objects.filter(user=user).count(), 1)

    def test_warning_when_teacher_has_no_subject(self):
        response = self.add_user(**{"profile-0-subjects": []}, _save="")
        response = self.client.get(response["Location"])
        self.assertContains(response, "оған пән тағайындалмаған")

    def test_assign_subject_on_change_page(self):
        teacher = User.objects.create_user(username="old", password="pass12345", is_staff=True)
        url = reverse("admin:auth_user_change", args=[teacher.pk])
        data = {
            "username": "old",
            "first_name": "",
            "last_name": "",
            "email": "",
            "is_active": "on",
            "is_staff": "on",
            "last_login_0": "",
            "last_login_1": "",
            "date_joined_0": "2026-10-01",
            "date_joined_1": "10:00:00",
            "profile-TOTAL_FORMS": "1",
            "profile-INITIAL_FORMS": "1",
            "profile-MIN_NUM_FORMS": "0",
            "profile-MAX_NUM_FORMS": "1",
            "profile-0-id": teacher.profile.pk,
            "profile-0-user": teacher.pk,
            "profile-0-group": "",
            "profile-0-subjects": [self.subject.pk],
            "_save": "Save",
        }
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(list(teacher.profile.subjects.all()), [self.subject])


class TeacherAdminAccessTests(TestCase):
    """Оқытушы (is_staff) admin-де тек өз пәнінің сессиялары мен топтарын басқарады."""

    def setUp(self):
        call_command("load_subjects", stdout=StringIO())
        self.mathematics = Subject.objects.get(code="mathematics")
        self.teacher = User.objects.create_user(
            username="teacher", password="pass12345", is_staff=True
        )
        self.teacher.profile.subjects.add(self.mathematics)
        self.client.force_login(self.teacher)
        self.math_group = StudyGroup.objects.create(name="МАТ-21", subject=self.mathematics, language="kk")
        self.inf_group = StudyGroup.objects.create(name="ИНФ-21", subject=informatics(), language="kk")
        now = timezone.now()
        self.inf_session = ExamSession.objects.create(
            title="Информатика сессиясы",
            subject=informatics(),
            opens_at=now,
            closes_at=now + timedelta(hours=2),
        )

    def session_data(self, subject, *groups):
        today = timezone.localdate().strftime("%Y-%m-%d")
        return {
            "title": "Математика сессиясы",
            "subject": subject.pk,
            "opens_at_0": today,
            "opens_at_1": "09:00:00",
            "closes_at_0": today,
            "closes_at_1": "18:00:00",
            "groups": [group.pk for group in groups],
            "show_answers": ExamSession.ShowAnswers.AFTER_FINISH,
            "is_active": "on",
            "_save": "Save",
        }

    def test_teacher_can_create_session_for_own_subject(self):
        self.assertEqual(self.client.get(reverse("admin:quiz_examsession_add")).status_code, 200)
        response = self.client.post(
            reverse("admin:quiz_examsession_add"), self.session_data(self.mathematics, self.math_group)
        )
        self.assertEqual(response.status_code, 302)
        session = ExamSession.objects.get(title="Математика сессиясы")
        self.assertEqual(list(session.groups.all()), [self.math_group])

    def test_teacher_cannot_create_session_for_other_subject(self):
        response = self.client.post(
            reverse("admin:quiz_examsession_add"), self.session_data(informatics())
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("subject", response.context["adminform"].form.errors)
        response = self.client.post(
            reverse("admin:quiz_examsession_add"),
            self.session_data(self.mathematics, self.inf_group),
        )
        self.assertIn("groups", response.context["adminform"].form.errors)
        self.assertFalse(ExamSession.objects.filter(title="Математика сессиясы").exists())

    def test_teacher_sees_only_own_subject(self):
        response = self.client.get(reverse("admin:quiz_examsession_changelist"))
        self.assertNotContains(response, "Информатика сессиясы")
        response = self.client.get(reverse("admin:accounts_studygroup_changelist"))
        self.assertContains(response, "МАТ-21")
        self.assertNotContains(response, "ИНФ-21")
        url = reverse("admin:quiz_examsession_change", args=[self.inf_session.pk])
        # Бөтен жазба — admin оны «жоқ» деп, басты бетке қайтарады
        self.assertEqual(self.client.get(url).status_code, 302)

    def test_teacher_can_create_group(self):
        response = self.client.post(
            reverse("admin:accounts_studygroup_add"),
            {"name": "МАТ-22", "subject": self.mathematics.pk, "language": "ru", "_save": "Save"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(StudyGroup.objects.get(name="МАТ-22").language, "ru")

    def test_group_language_is_required(self):
        response = self.client.post(
            reverse("admin:accounts_studygroup_add"),
            {"name": "МАТ-23", "subject": self.mathematics.pk, "_save": "Save"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("language", response.context["adminform"].form.errors)
        self.assertFalse(StudyGroup.objects.filter(name="МАТ-23").exists())

    def test_teacher_cannot_manage_users_or_delete(self):
        forbidden = [
            reverse("admin:auth_user_changelist"),
            reverse("admin:auth_user_add"),
            reverse("admin:quiz_question_changelist"),
            reverse("admin:accounts_studygroup_delete", args=[self.math_group.pk]),
        ]
        for url in forbidden:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 403)

    def test_admin_index_lists_sessions_and_groups(self):
        response = self.client.get(reverse("admin:index"))
        self.assertContains(response, reverse("admin:quiz_examsession_changelist"))
        self.assertContains(response, reverse("admin:accounts_studygroup_changelist"))
        self.assertNotContains(response, reverse("admin:auth_user_changelist"))

    def test_teacher_cannot_import_students(self):
        response = self.client.get(reverse("admin:accounts_studygroup_changelist"))
        self.assertNotContains(response, reverse("admin:accounts_studygroup_import"))
        self.assertNotContains(response, "regenerate_group_codes")
        for url_name in ["accounts_studygroup_import", "accounts_studygroup_import_template"]:
            with self.subTest(url_name=url_name):
                self.assertEqual(self.client.get(reverse(f"admin:{url_name}")).status_code, 403)


# ---------- Студенттерді Excel-ден импорттау, ИИН + код арқылы кіру ----------

# Бақылау цифры дұрыс ИИН-дер
IIN_AIGERIM = "050312650125"
IIN_DANA = "041122500343"
IIN_ERLAN = "980701400565"


def students_file(rows, name="students.xlsx"):
    """Студенттер файлы: бірінші парақ, 1-жол — тақырыптар."""
    workbook = Workbook()
    workbook.active.append(["Аты-жөні", "Топ", "ИИН"])
    for row in rows:
        workbook.active.append(list(row))
    output = BytesIO()
    workbook.save(output)
    return SimpleUploadedFile(name, output.getvalue())


def read_codes(response):
    """Жүктелген кодтар файлы: {ИИН: (аты-жөні, топ, код)}."""
    workbook = load_workbook(BytesIO(response.content))
    rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
    return {iin: (name, group, code) for name, group, iin, code in rows}


class IinTests(TestCase):
    def test_valid_iins(self):
        for iin in [IIN_AIGERIM, IIN_DANA, IIN_ERLAN]:
            with self.subTest(iin=iin):
                self.assertTrue(iin_is_valid(iin))

    def test_invalid_iins(self):
        for iin in ["050312650124", "05031265012", "0503126501255", "05031265012a", ""]:
            with self.subTest(iin=iin):
                self.assertFalse(iin_is_valid(iin))

    def test_normalize_iin_restores_leading_zero(self):
        # Excel ИИН-ді сан деп сақтаса, бас нөл жоғалады: 050312650125 → 50312650125
        self.assertEqual(normalize_iin(50312650125), IIN_AIGERIM)
        self.assertEqual(normalize_iin("0503 1265 0125"), IIN_AIGERIM)

    def test_split_name(self):
        self.assertEqual(split_name("Сапарова Айгерім Нұрланқызы"), ("Сапарова", "Айгерім Нұрланқызы"))
        self.assertEqual(split_name("Сапарова"), ("Сапарова", ""))


class StudentImportTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(username="admin", password="pass12345")
        self.client.force_login(self.admin)
        self.inf = StudyGroup.objects.create(name="ИНФ-21", subject=informatics(), language="kk")
        self.inf2 = StudyGroup.objects.create(name="ИНФ-22", subject=informatics(), language="ru")

    def upload(self, rows):
        return self.client.post(
            reverse("admin:accounts_studygroup_import"), {"file": students_file(rows)}
        )

    def test_page_opens_from_group_list(self):
        response = self.client.get(reverse("admin:accounts_studygroup_changelist"))
        self.assertContains(response, reverse("admin:accounts_studygroup_import"))
        response = self.client.get(reverse("admin:accounts_studygroup_import"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse("admin:accounts_studygroup_import_template"))

    def test_template_lists_groups(self):
        response = self.client.get(reverse("admin:accounts_studygroup_import_template"))
        self.assertIn("students-template-", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content))
        groups = [row[0] for row in workbook.worksheets[1].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(groups, ["ИНФ-21", "ИНФ-22"])

    def test_import_creates_students_and_returns_codes(self):
        response = self.upload(
            [
                ["Сапарова Айгерім Нұрланқызы", "ИНФ-21", IIN_AIGERIM],
                ["  Әлиева   Дана ", "инф-22", int(IIN_DANA)],
            ]
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("student-codes-", response["Content-Disposition"])
        codes = read_codes(response)
        self.assertEqual(set(codes), {IIN_AIGERIM, IIN_DANA})

        aigerim = User.objects.get(username=IIN_AIGERIM)
        self.assertEqual((aigerim.last_name, aigerim.first_name), ("Сапарова", "Айгерім Нұрланқызы"))
        self.assertEqual(aigerim.profile.group, self.inf)
        self.assertFalse(aigerim.is_staff)
        dana = User.objects.get(username=IIN_DANA)
        self.assertEqual((dana.last_name, dana.first_name), ("Әлиева", "Дана"))
        self.assertEqual(dana.profile.group, self.inf2)

        name, group, code = codes[IIN_AIGERIM]
        self.assertEqual((name, group), ("Сапарова Айгерім Нұрланқызы", "ИНФ-21"))
        self.assertRegex(code, r"^\d{8}$")
        # Код ашық түрде сақталмайды
        self.assertNotIn(code, aigerim.password)

        # Студент ИИН мен кодпен кіреді (бос орынмен терсе де)
        self.client.logout()
        response = self.client.post(
            reverse("accounts:login"), {"username": "0503 1265 0125", "password": code}
        )
        self.assertRedirects(response, reverse("accounts:after_login"), fetch_redirect_response=False)

    def test_success_message_on_next_page(self):
        self.upload([["Сапарова Айгерім", "ИНФ-21", IIN_AIGERIM]])
        response = self.client.get(reverse("admin:accounts_studygroup_changelist"))
        self.assertContains(response, "жаңасы — 1, жаңартылғаны — 0")

    def test_errors_save_nothing(self):
        User.objects.create_user(username=IIN_ERLAN, password="pass12345", is_staff=True)
        response = self.upload(
            [
                ["Сапарова Айгерім", "ИНФ-21", IIN_AIGERIM],
                ["", "ИНФ-99", "050312650124"],
                ["Әлиева Дана", "ИНФ-21", IIN_AIGERIM],
                ["Ерлан Серікұлы", "ИНФ-21", IIN_ERLAN],
            ]
        )
        self.assertEqual(response.status_code, 200)
        errors = response.context["result"]["errors"]
        self.assertEqual(
            errors,
            [
                "3-жол, «Аты-жөні»: бос болмауы керек.",
                "3-жол, «Топ»: «ИНФ-99» тобы жоқ: алдымен оны admin-де құрыңыз.",
                "3-жол, «ИИН»: 12 цифрдан тұрып, бақылау цифры дұрыс болуы керек.",
                "4-жол, «ИИН»: бұл ИИН 2-жолда бар.",
                "5-жол, «ИИН»: бұл — қызметкердің логині, оны өзгертуге болмайды.",
            ],
        )
        self.assertContains(response, "ештеңе сақталмады")
        self.assertFalse(User.objects.filter(username=IIN_AIGERIM).exists())

    def test_reimport_updates_name_and_group_but_keeps_code(self):
        response = self.upload([["Сапарова Айгерім", "ИНФ-21", IIN_AIGERIM]])
        code = read_codes(response)[IIN_AIGERIM][2]
        response = self.upload(
            [["Сапарова Айгерім Нұрланқызы", "ИНФ-22", IIN_AIGERIM], ["Әлиева Дана", "ИНФ-21", IIN_DANA]]
        )
        codes = read_codes(response)
        self.assertEqual(codes[IIN_AIGERIM][2], "бұрынғы код")
        self.assertRegex(codes[IIN_DANA][2], r"^\d{8}$")
        aigerim = User.objects.get(username=IIN_AIGERIM)
        self.assertEqual(aigerim.first_name, "Айгерім Нұрланқызы")
        self.assertEqual(aigerim.profile.group, self.inf2)
        self.assertTrue(aigerim.check_password(code))
        self.assertEqual(User.objects.filter(username=IIN_AIGERIM).count(), 1)

    def test_file_checks(self):
        response = self.client.post(
            reverse("admin:accounts_studygroup_import"),
            {"file": SimpleUploadedFile("students.csv", b"a;b")},
        )
        self.assertIn("file", response.context["form"].errors)
        response = self.upload([])
        self.assertEqual(response.context["result"]["errors"], ["Файлда студент жоқ."])

    def test_first_login_rehashes_code_with_default_hasher(self):
        response = self.upload([["Сапарова Айгерім", "ИНФ-21", IIN_AIGERIM]])
        code = read_codes(response)[IIN_AIGERIM][2]
        user = User.objects.get(username=IIN_AIGERIM)
        self.assertTrue(user.password.startswith("pbkdf2_access_code$"))
        self.client.logout()
        self.client.post(reverse("accounts:login"), {"username": IIN_AIGERIM, "password": code})
        user.refresh_from_db()
        self.assertTrue(user.password.startswith("pbkdf2_sha256$"))

    def test_regenerate_codes_for_group(self):
        response = self.upload([["Сапарова Айгерім", "ИНФ-21", IIN_AIGERIM]])
        old_code = read_codes(response)[IIN_AIGERIM][2]
        response = self.client.post(
            reverse("admin:accounts_studygroup_changelist"),
            {"action": "regenerate_group_codes", "_selected_action": [self.inf.pk]},
        )
        new_code = read_codes(response)[IIN_AIGERIM][2]
        user = User.objects.get(username=IIN_AIGERIM)
        self.assertTrue(user.check_password(new_code))
        if new_code != old_code:
            self.assertFalse(user.check_password(old_code))

    def test_regenerate_codes_skips_staff(self):
        self.upload([["Сапарова Айгерім", "ИНФ-21", IIN_AIGERIM]])
        student = User.objects.get(username=IIN_AIGERIM)
        response = self.client.post(
            reverse("admin:auth_user_changelist"),
            {"action": "regenerate_user_codes", "_selected_action": [student.pk, self.admin.pk]},
        )
        self.assertEqual(set(read_codes(response)), {IIN_AIGERIM})
        self.admin.refresh_from_db()
        self.assertTrue(self.admin.check_password("pass12345"))

    def test_page_in_russian(self):
        self.client.cookies["django_language"] = "ru"
        response = self.client.get(reverse("admin:accounts_studygroup_import"))
        self.assertContains(response, "Импорт студентов")
        response = self.upload([["", "ИНФ-21", IIN_AIGERIM]])
        self.assertContains(response, "строка 2, «Имя и фамилия»: не должно быть пустым.")


class LoginFormTests(TestCase):
    def test_login_labels(self):
        response = self.client.get(reverse("accounts:login"))
        self.assertContains(response, "ИИН немесе логин")
        self.assertContains(response, "Код немесе құпия сөз")
        self.assertContains(response, "Кодты әкімші береді")

    def test_teacher_logs_in_with_username(self):
        User.objects.create_user(username="teacher", password="pass12345", is_staff=True)
        response = self.client.post(
            reverse("accounts:login"), {"username": "teacher", "password": "pass12345"}
        )
        self.assertEqual(response.status_code, 302)


class StudentsNavLinkTests(TestCase):
    """Навбардағы «Студенттер» — тек әкімшіге, admin-дегі топтар тізіміне апарады."""

    def test_superuser_sees_link(self):
        admin_user = User.objects.create_superuser(username="admin", password="pass12345")
        self.client.force_login(admin_user)
        response = self.client.get(reverse("quiz:teacher_questions"))
        self.assertContains(response, reverse("admin:accounts_studygroup_changelist"))
        self.assertContains(response, ">Студенттер</a>")

    def test_teacher_and_student_do_not_see_link(self):
        teacher = User.objects.create_user(username="teacher", password="pass12345", is_staff=True)
        teacher.profile.subjects.add(informatics())
        student = User.objects.create_user(username="student", password="pass12345")
        for user, url in [(teacher, "quiz:teacher_questions"), (student, "quiz:dashboard")]:
            with self.subTest(user=user.username):
                self.client.force_login(user)
                response = self.client.get(reverse(url))
                self.assertNotContains(response, reverse("admin:accounts_studygroup_changelist"))


class AddStudentTests(TestCase):
    """Admin-дегі «Студент қосу»: бір студент, код экранда бір рет көрсетіледі."""

    def setUp(self):
        self.admin = User.objects.create_superuser(username="admin", password="pass12345")
        self.client.force_login(self.admin)
        self.group = StudyGroup.objects.create(name="ИНФ-21", subject=informatics(), language="kk")
        self.url = reverse("admin:accounts_studygroup_add_student")

    def add(self, **overrides):
        data = {"full_name": "Сапарова Айгерім Нұрланқызы", "group": self.group.pk, "iin": IIN_AIGERIM}
        data.update(overrides)
        return self.client.post(self.url, data)

    def test_button_on_group_list(self):
        response = self.client.get(reverse("admin:accounts_studygroup_changelist"))
        self.assertContains(response, self.url)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ИНФ-21 — Информатика (Қазақша)")

    def test_add_student_shows_code_once(self):
        response = self.add(iin="0503 1265 0125", full_name="  Сапарова   Айгерім Нұрланқызы ")
        self.assertEqual(response.status_code, 200)
        self.assertIn("no-cache", response["Cache-Control"])
        student = User.objects.get(username=IIN_AIGERIM)
        self.assertEqual((student.last_name, student.first_name), ("Сапарова", "Айгерім Нұрланқызы"))
        self.assertEqual(student.profile.group, self.group)
        self.assertFalse(student.is_staff)

        code = response.context["created"]["code"]
        self.assertRegex(code, r"^\d{8}$")
        self.assertContains(response, f'<strong id="access-code"')
        self.assertContains(response, code)
        self.assertTrue(student.check_password(code))
        # Форма тазаланып, келесі студентті қосуға дайын
        self.assertFalse(response.context["form"].is_bound)

        # Бетті қайта ашқанда код көрінбейді
        self.assertNotContains(self.client.get(self.url), code)

    def test_errors(self):
        cases = {
            "full_name": {"full_name": "   "},
            "group": {"group": ""},
            "iin": {"iin": "050312650124"},
        }
        for field, overrides in cases.items():
            with self.subTest(field=field):
                response = self.add(**overrides)
                self.assertIn(field, response.context["form"].errors)
        self.assertFalse(User.objects.filter(username=IIN_AIGERIM).exists())

    def test_existing_iin_links_to_student(self):
        self.add()
        student = User.objects.get(username=IIN_AIGERIM)
        response = self.add(full_name="Басқа Адам")
        self.assertContains(response, "Бұл ИИН-мен студент бұрыннан бар.")
        self.assertContains(response, reverse("admin:auth_user_change", args=[student.pk]))
        self.assertEqual(User.objects.filter(username=IIN_AIGERIM).count(), 1)
        student.refresh_from_db()
        self.assertEqual(student.last_name, "Сапарова")

    def test_staff_iin_is_rejected(self):
        User.objects.create_user(username=IIN_ERLAN, password="pass12345", is_staff=True)
        response = self.add(iin=IIN_ERLAN)
        self.assertContains(response, "қызметкердің логині")

    def test_teacher_cannot_add_student(self):
        teacher = User.objects.create_user(username="teacher", password="pass12345", is_staff=True)
        teacher.profile.subjects.add(informatics())
        self.client.force_login(teacher)
        response = self.client.get(reverse("admin:accounts_studygroup_changelist"))
        self.assertNotContains(response, self.url)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertEqual(self.add().status_code, 403)
        self.assertFalse(User.objects.filter(username=IIN_AIGERIM).exists())

    def test_page_in_russian(self):
        self.client.cookies["django_language"] = "ru"
        response = self.add()
        self.assertContains(response, "Студент добавлен.")
        self.assertContains(response, "Добавить студента")

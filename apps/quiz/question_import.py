"""
Сұрақтарды Excel (.xlsx) файлынан импорттау және бос үлгі файлын жасау.

Файлда екі парақ:
- «Сұрақтар» — бір жол = бір сұрақ: тақырыпша, деңгей, тіл, контекст кілті,
  мәтін, код, A–D нұсқалары, дұрыс жауап әрпі, сурет;
- «Контексттер» (міндетті емес) — кілт, атауы, тіл, мәтін, код, сурет.
Бағандар тақырыбы бойынша емес, реті бойынша оқылады (1-жол — тақырыптар).

Суреттер Excel-ге ұяшықтың үстіне қойылады: сурет сол жақ жоғарғы бұрышы тұрған
ұяшыққа тиесілі. Сұрақ пен контекст суреті — «Сурет» бағанында, жауап суреті —
сол жауаптың (A–D) ұяшығында.

Тексеру «бәрі немесе ештеңе»: бір қате болса — ештеңе сақталмайды, қателер
жол нөмірімен қайтарылады. Пәнде бұрыннан бар сұрақ (мәтіні, тілі, тақырыпшасы
бірдей) қайта қосылмайды.
"""

import uuid
import zipfile
from io import BytesIO

from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db import transaction
from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from .constants import ANSWERS_PER_QUESTION, IMPORT_MAX_ROWS, MAX_IMAGE_MB, QUESTIONS_PER_CONTEXT
from .models import Answer, Context, Language, Level, Question, Subtopic
from .services import ANSWER_LETTERS

# Excel (.xlsx) файлының MIME түрі
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Парақ атаулары аударылмайды: импорт оларды осы атаумен іздейді
QUESTIONS_SHEET = "Сұрақтар"
CONTEXTS_SHEET = "Контексттер"

# Бағандар реті (0-ден): «Сұрақтар» парағы
SUBTOPIC, LEVEL, LANGUAGE, CONTEXT, TEXT, CODE = range(6)
FIRST_ANSWER = 6  # A–D нұсқалары: 6, 7, 8, 9
CORRECT = FIRST_ANSWER + ANSWERS_PER_QUESTION
IMAGE = CORRECT + 1
QUESTION_WIDTH = IMAGE + 1

# «Контексттер» парағы
CONTEXT_KEY, CONTEXT_TITLE, CONTEXT_LANGUAGE, CONTEXT_TEXT, CONTEXT_CODE, CONTEXT_IMAGE = range(6)
CONTEXT_WIDTH = 6

# Сурет пішімі → файл кеңейтімі. openpyxl webp т.б. пішімдерді PNG-ге айналдырып береді,
# тек gif өзгеріссіз қалады — ол рұқсат етілмейді
IMAGE_EXTENSIONS = {"jpeg": "jpg", "png": "png", "webp": "png", "bmp": "png", "tiff": "png"}

# Excel 365-тің «ұяшықтың ішіне» қойылған суреті openpyxl-ге көрінбейді, ұяшықта осы қате тұрады
IN_CELL_IMAGE = "#VALUE!"

# Кириллицаның А, В, С, Д әріптері латынша сияқты көрінеді — оларды латыншаға ауыстырамыз
CYRILLIC_TO_LATIN = str.maketrans({"А": "A", "В": "B", "С": "C", "Д": "D"})

# Тіл бағанында рұқсат етілген жазулар
LANGUAGE_ALIASES = {
    "kk": Language.KK,
    "kz": Language.KK,
    "қаз": Language.KK,
    "қазақша": Language.KK,
    "ru": Language.RU,
    "рус": Language.RU,
    "орыс": Language.RU,
    "орысша": Language.RU,
}

# Excel-де тізімнен таңдау (Деректерді тексеру) қолданылатын жол саны
VALIDATION_ROWS = IMPORT_MAX_ROWS + 1


def question_headers():
    """«Сұрақтар» парағының баған тақырыптары (интерфейс тілінде)."""
    return [
        _("Тақырыпша (01–20)"),
        _("Деңгей (A/B/C)"),
        _("Тіл (kk/ru)"),
        _("Контекст кілті"),
        _("Сұрақ мәтіні"),
        _("Код"),
        *ANSWER_LETTERS[:ANSWERS_PER_QUESTION],
        _("Дұрыс жауап (A–D)"),
        _("Сурет"),
    ]


def context_headers():
    """«Контексттер» парағының баған тақырыптары."""
    return [_("Кілт"), _("Атауы"), _("Тіл (kk/ru)"), _("Мәтіні"), _("Код"), _("Сурет")]


# ---------- Ұяшықтарды оқу ----------


def cell_text(value):
    """Ұяшық мәні мәтін ретінде: бос — "", 3.0 сияқты бүтін сан — "3"."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\r\n", "\n")


def code_text(value):
    """Код пен жауап нұсқасы: шегіністер сақталады, тек бос орындар болса — ""."""
    text = cell_text(value)
    return text if text.strip() else ""


def parse_number(value):
    """Тақырыпша нөмірі: 3, "03", 3.0 → 3; сан болмаса — None."""
    text = cell_text(value).strip()
    return int(text) if text.isdigit() else None


def parse_letter(value):
    """Деңгей мен дұрыс жауап әрпі: бас әріп, кириллица латыншаға ауыстырылады."""
    return cell_text(value).strip().upper().translate(CYRILLIC_TO_LATIN)


def parse_language(value):
    """Тіл: kk / ru (қаз, орыс т.б. да болады); танылмаса — None."""
    return LANGUAGE_ALIASES.get(cell_text(value).strip().lower())


class TooManyRows(Exception):
    """Парақта IMPORT_MAX_ROWS-тан көп жол бар."""


def read_rows(sheet, width):
    """
    Парақтың деректер жолдары (1-жол — тақырыптар), бос жолдарсыз:
    [(жол нөмірі, [ұяшық мәндері]), ...]. Әр жолда дәл width ұяшық.
    """
    rows = []
    cells = sheet.iter_rows(min_row=2, max_col=width, values_only=True)
    for number, values in enumerate(cells, start=2):
        values = list(values) + [None] * (width - len(values))
        if all(cell_text(value).strip() == "" for value in values):
            continue
        if len(rows) >= IMPORT_MAX_ROWS:
            raise TooManyRows
        rows.append((number, values))
    return rows


# ---------- Суреттер ----------


def sheet_images(sheet):
    """
    Парақтағы суреттер: {(жол нөмірі, баған индексі 0-ден): [сурет, ...]}.
    Ұяшыққа байланбаған сурет (абсолют орны бар) — (None, None) кілтінде.
    openpyxl-де суреттерді оқитын ашық API жоқ, сондықтан sheet._images қолданылады.
    """
    images = {}
    for image in sheet._images:
        start = getattr(image.anchor, "_from", None)
        place = (start.row + 1, start.col) if start is not None else (None, None)
        images.setdefault(place, []).append(image)
    return images


def image_content(image):
    """Сурет файлы: ((байттар, кеңейтім), None) не пішімі/көлемі жарамсыз болса — (None, қате)."""
    extension = IMAGE_EXTENSIONS.get(image.format)
    if extension is None:
        return None, _("сурет пішімі jpg, png немесе webp болуы керек.")
    data = image._data()
    if len(data) > MAX_IMAGE_MB * 1024 * 1024:
        return None, _("сурет тым үлкен: ең көбі %(size)s МБ.") % {"size": MAX_IMAGE_MB}
    return (data, extension), None


def read_images(sheet, sheet_title, rows, image_columns, headers, add_error):
    """
    Парақтағы суреттерді тексеріп оқиды: {(жол, баған): (байттар, кеңейтім)}.
    Сурет деректері бар жолда, рұқсат етілген бағанда тұруы керек; бір ұяшықта — бір сурет.
    """
    data_rows = {number for number, _cells in rows}
    allowed = ", ".join(f"«{headers[column]}»" for column in image_columns)
    contents = {}
    for (row, column), images in sheet_images(sheet).items():
        if row is None:
            add_error(sheet_title, "—", _("Сурет"), _("сурет ұяшыққа байланбаған."))
            continue
        name = headers[column] if column < len(headers) else get_column_letter(column + 1)
        if row not in data_rows:
            add_error(sheet_title, row, name, _("суреті бар жолда мәлімет жоқ."))
        elif column not in image_columns:
            add_error(
                sheet_title,
                row,
                name,
                _("сурет мына бағандардың бірінде тұруы керек: %(columns)s.")
                % {"columns": allowed},
            )
        elif len(images) > 1:
            add_error(sheet_title, row, name, _("бір ұяшықта бір ғана сурет болуы керек."))
        else:
            content, error = image_content(images[0])
            if error:
                add_error(sheet_title, row, name, error)
            else:
                contents[(row, column)] = content
    return contents


def check_in_cell_images(sheet_title, rows, columns, headers, add_error):
    """
    Сурет қойылатын бағандарда (сұрақ/контекст суреті, A–D жауаптары) ұяшықтың
    ішіне қойылған сурет болса — қате: импорт оны көрмейді, ал жауап ұяшығында
    «#VALUE!» мәтін болып сақталып қалар еді.
    """
    for row, cells in rows:
        for column in columns:
            if cell_text(cells[column]).strip() == IN_CELL_IMAGE:
                add_error(
                    sheet_title,
                    row,
                    headers[column],
                    _(
                        "сурет ұяшықтың ішіне қойылған — оны ұяшықтың үстіне қойыңыз "
                        "(Кірістіру → Суреттер → Ұяшықтардың үстіне орналастыру)."
                    ),
                )


def image_file(content):
    """Сақталатын сурет файлы, кездейсоқ атпен (uuid)."""
    data, extension = content
    return ContentFile(data, name=f"{uuid.uuid4().hex}.{extension}")


# ---------- Импорт ----------


def new_result():
    return {"errors": [], "questions": 0, "contexts": 0, "skipped": []}


def normalize_text(text):
    """Қайталануды салыстыру үшін: жол соңы бірдей, шеттегі бос орындарсыз."""
    return text.replace("\r\n", "\n").strip()


def import_questions(subject, file):
    """
    Excel файлындағы сұрақтар мен контексттерді пәнге қосады.

    Қайтарады: {"errors": [...], "questions": қосылған сұрақтар саны,
    "contexts": қосылған контексттер саны, "skipped": өткізілген жолдар нөмірі}.
    errors бос болмаса — дерекқорға ештеңе жазылмаған.
    """
    result = new_result()
    errors = result["errors"]

    try:
        # read_only режимінде суреттер оқылмайды, сондықтан файл толық ашылады
        workbook = load_workbook(file, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError):
        errors.append(_("Файл оқылмады: ол Excel (.xlsx) файлы болуы керек."))
        return result

    try:
        if QUESTIONS_SHEET in workbook.sheetnames:
            question_sheet = workbook[QUESTIONS_SHEET]
        else:
            # Өз файлы болса: «Сұрақтар» парағы жоқ — бірінші парақ
            question_sheet = workbook.worksheets[0]
        question_rows = read_rows(question_sheet, QUESTION_WIDTH)
        context_sheet = None
        context_rows = []
        if CONTEXTS_SHEET in workbook.sheetnames and question_sheet.title != CONTEXTS_SHEET:
            context_sheet = workbook[CONTEXTS_SHEET]
            context_rows = read_rows(context_sheet, CONTEXT_WIDTH)
    except TooManyRows:
        errors.append(
            _("Файлда жол тым көп: бір парақта ең көбі %(count)s жол.")
            % {"count": IMPORT_MAX_ROWS}
        )
        return result
    finally:
        workbook.close()

    def add_error(sheet, row, column, message):
        errors.append(
            _("«%(sheet)s» парағы, %(row)s-жол, «%(column)s»: %(message)s")
            % {"sheet": sheet, "row": row, "column": column, "message": message}
        )

    question_images = read_images(
        question_sheet,
        QUESTIONS_SHEET,
        question_rows,
        [IMAGE, *range(FIRST_ANSWER, CORRECT)],
        question_headers(),
        add_error,
    )
    check_in_cell_images(
        QUESTIONS_SHEET,
        question_rows,
        [IMAGE, *range(FIRST_ANSWER, CORRECT)],
        question_headers(),
        add_error,
    )
    context_images = {}
    if context_sheet is not None:
        context_images = read_images(
            context_sheet, CONTEXTS_SHEET, context_rows, [CONTEXT_IMAGE], context_headers(), add_error
        )
        check_in_cell_images(
            CONTEXTS_SHEET, context_rows, [CONTEXT_IMAGE], context_headers(), add_error
        )

    contexts = read_contexts(subject, context_rows, context_images, add_error)
    questions = read_questions(subject, question_rows, question_images, contexts, add_error, result)

    # Контекстте 5-тен көп белсенді сұрақ болмайды (бұрынғыларымен бірге)
    for key, data in contexts.items():
        total = data["existing_count"] + data["new_count"]
        if total > QUESTIONS_PER_CONTEXT:
            add_error(
                CONTEXTS_SHEET,
                data["row"],
                context_headers()[CONTEXT_KEY],
                _("«%(key)s» контекстіне %(total)s сұрақ байланады, ең көбі %(max)s.")
                % {"key": key, "total": total, "max": QUESTIONS_PER_CONTEXT},
            )

    if not errors and not questions and not result["skipped"] and not contexts:
        errors.append(_("Файлда сұрақ жоқ."))
    if errors:
        result["skipped"] = []
        return result

    save_import(subject, contexts, questions, result)
    return result


def read_contexts(subject, rows, images, add_error):
    """
    «Контексттер» парағы: {кілт: {...}}. Пәнде атауы, тілі, мәтіні бірдей контекст
    бұрыннан бар болса — жаңасы жасалмайды, сұрақтар соған байланады.
    """
    headers = context_headers()
    existing = {}
    for quiz_context in Context.objects.filter(subject=subject):
        key = (
            normalize_text(quiz_context.title),
            quiz_context.language,
            normalize_text(quiz_context.text),
        )
        existing.setdefault(key, quiz_context)

    contexts = {}
    for row, cells in rows:
        key = cell_text(cells[CONTEXT_KEY]).strip()
        title = cell_text(cells[CONTEXT_TITLE]).strip()
        language = parse_language(cells[CONTEXT_LANGUAGE])
        text = cell_text(cells[CONTEXT_TEXT]).strip()

        if not key:
            add_error(CONTEXTS_SHEET, row, headers[CONTEXT_KEY], _("бос болмауы керек."))
            continue
        if key in contexts:
            add_error(
                CONTEXTS_SHEET,
                row,
                headers[CONTEXT_KEY],
                _("«%(key)s» кілті %(row)s-жолда бар.")
                % {"key": key, "row": contexts[key]["row"]},
            )
            continue
        if not title:
            add_error(CONTEXTS_SHEET, row, headers[CONTEXT_TITLE], _("бос болмауы керек."))
        if language is None:
            add_error(CONTEXTS_SHEET, row, headers[CONTEXT_LANGUAGE], _("kk немесе ru болуы керек."))
        if not text:
            add_error(CONTEXTS_SHEET, row, headers[CONTEXT_TEXT], _("бос болмауы керек."))

        found = existing.get((normalize_text(title), language, normalize_text(text)))
        contexts[key] = {
            "row": row,
            "title": title,
            "language": language,
            "text": text,
            "code": code_text(cells[CONTEXT_CODE]),
            "image": images.get((row, CONTEXT_IMAGE)),
            "instance": found,
            "existing_count": found.active_questions_count() if found else 0,
            "new_count": 0,
        }
    return contexts


def read_questions(subject, rows, images, contexts, add_error, result):
    """
    «Сұрақтар» парағы: сақталатын сұрақтар тізімі. Қателер add_error арқылы
    жиналады, бұрыннан бар сұрақтардың жол нөмірі result["skipped"]-ке жазылады.
    """
    headers = question_headers()
    subtopics = {
        subtopic.number: subtopic
        for subtopic in Subtopic.objects.filter(topic__subject=subject)
    }
    seen = {
        (normalize_text(text), language, subtopic_id)
        for text, language, subtopic_id in Question.objects.filter(
            subtopic__topic__subject=subject
        ).values_list("text", "language", "subtopic_id")
    }
    sheet = QUESTIONS_SHEET
    letters = ANSWER_LETTERS[:ANSWERS_PER_QUESTION]

    questions = []
    for row, cells in rows:
        errors_before = len(result["errors"])

        subtopic = subtopics.get(parse_number(cells[SUBTOPIC]))
        if subtopic is None:
            add_error(
                sheet,
                row,
                headers[SUBTOPIC],
                _("пәнде мұндай тақырыпша жоқ (01–%(max)s).") % {"max": len(subtopics)},
            )
        level = parse_letter(cells[LEVEL])
        if level not in Level.values:
            add_error(sheet, row, headers[LEVEL], _("A, B немесе C болуы керек."))
        language = parse_language(cells[LANGUAGE])
        if language is None:
            add_error(sheet, row, headers[LANGUAGE], _("kk немесе ru болуы керек."))
        text = cell_text(cells[TEXT]).strip()
        if not text:
            add_error(sheet, row, headers[TEXT], _("бос болмауы керек."))

        answers = [code_text(cells[FIRST_ANSWER + index]) for index in range(len(letters))]
        answer_images = [images.get((row, FIRST_ANSWER + index)) for index in range(len(letters))]
        for letter, answer, answer_image in zip(letters, answers, answer_images):
            if not answer and not answer_image:
                add_error(sheet, row, letter, _("жауап нұсқасында мәтін не сурет болуы керек."))
        correct = parse_letter(cells[CORRECT])
        if correct not in letters:
            add_error(
                sheet,
                row,
                headers[CORRECT],
                _("%(letters)s әріптерінің бірі болуы керек.") % {"letters": ", ".join(letters)},
            )

        context_key = cell_text(cells[CONTEXT]).strip()
        context = None
        if context_key:
            context = contexts.get(context_key)
            if context is None:
                add_error(
                    sheet,
                    row,
                    headers[CONTEXT],
                    _("«%(key)s» кілті «%(contexts)s» парағында жоқ.")
                    % {"key": context_key, "contexts": CONTEXTS_SHEET},
                )
            elif language and context["language"] and context["language"] != language:
                add_error(
                    sheet,
                    row,
                    headers[CONTEXT],
                    _("контекст пен сұрақтың тілі бірдей болуы керек."),
                )

        if len(result["errors"]) > errors_before:
            continue

        # Пәнде (немесе файлдың жоғарғы жолдарында) бұрыннан бар сұрақ — өткізіледі
        duplicate_key = (normalize_text(text), language, subtopic.pk)
        if duplicate_key in seen:
            result["skipped"].append(row)
            continue
        seen.add(duplicate_key)

        if context is not None:
            context["new_count"] += 1
        questions.append(
            {
                "subtopic": subtopic,
                "level": level,
                "language": language,
                "context_key": context_key,
                "text": text,
                "code": code_text(cells[CODE]),
                "answers": answers,
                "answer_images": answer_images,
                "image": images.get((row, IMAGE)),
                "correct_index": letters.index(correct),
            }
        )
    return questions


def save_import(subject, contexts, questions, result):
    """
    Тексерілген контексттер мен сұрақтарды бір транзакцияда сақтайды.
    Сурет файлдары дискіге бірден жазылады, ал транзакция кері қайтарылса
    (сақтау ортасында қате шықса) олар дерекқорда жоқ болып қалады — сондықтан
    жазылған файлдар өшіріледі, қате әрі қарай жіберіледі.
    """
    saved_files = []

    def remember(instance):
        if instance.image:
            saved_files.append(instance.image.name)
        return instance

    try:
        with transaction.atomic():
            save_objects(subject, contexts, questions, result, remember)
    except Exception:
        for name in saved_files:
            default_storage.delete(name)
        raise


def save_objects(subject, contexts, questions, result, remember):
    """save_import-тің ішкі бөлігі: жазбаларды жасайды, әр жазбаны remember()-ге береді."""
    for data in contexts.values():
        if data["instance"] is None:
            data["instance"] = remember(
                Context.objects.create(
                    subject=subject,
                    language=data["language"],
                    title=data["title"],
                    text=data["text"],
                    code=data["code"],
                    image=image_file(data["image"]) if data["image"] else "",
                )
            )
            result["contexts"] += 1

    for data in questions:
        context = contexts[data["context_key"]]["instance"] if data["context_key"] else None
        question = remember(
            Question.objects.create(
                subtopic=data["subtopic"],
                context=context,
                language=data["language"],
                text=data["text"],
                code=data["code"],
                level=data["level"],
                image=image_file(data["image"]) if data["image"] else "",
            )
        )
        for index, (text, image) in enumerate(zip(data["answers"], data["answer_images"])):
            remember(
                Answer.objects.create(
                    question=question,
                    text=text,
                    image=image_file(image) if image else "",
                    is_correct=index == data["correct_index"],
                )
            )
        result["questions"] += 1


def result_message(result):
    """Сәтті импорттан кейінгі хабарлама."""
    message = _("Қосылды: сұрақтар — %(questions)s, контексттер — %(contexts)s.") % {
        "questions": result["questions"],
        "contexts": result["contexts"],
    }
    if result["skipped"]:
        rows = ", ".join(str(row) for row in result["skipped"])
        message += " " + _("Банкте бұрыннан бар, өткізілді — жолдар: %(rows)s.") % {"rows": rows}
    return message


# ---------- Үлгі файл ----------


def fill_sheet(sheet, headers, widths):
    """Тақырып жолы (қалың), бағандар ені, тақырып жолы бекітіледі."""
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = "A2"


def add_choice(sheet, column, choices):
    """Бағанда тек тізімдегі мән таңдалады (Excel-дің «Деректерді тексеру»)."""
    validation = DataValidation(
        type="list", formula1='"{}"'.format(",".join(choices)), allow_blank=True
    )
    sheet.add_data_validation(validation)
    validation.add(f"{column}2:{column}{VALIDATION_ROWS}")


def import_template(subject):
    """Бос үлгі файл (.xlsx, bytes): екі парақ, нұсқаулық және пәннің тақырыпшалары."""
    workbook = Workbook()
    letters = ANSWER_LETTERS[:ANSWERS_PER_QUESTION]

    questions = workbook.active
    questions.title = QUESTIONS_SHEET
    fill_sheet(questions, question_headers(), [12, 10, 9, 12, 50, 30, 20, 20, 20, 20, 12, 30])
    add_choice(questions, "B", Level.values)
    add_choice(questions, "C", Language.values)
    add_choice(questions, chr(ord("A") + CORRECT), letters)

    contexts = workbook.create_sheet(CONTEXTS_SHEET)
    fill_sheet(contexts, context_headers(), [10, 30, 9, 60, 30, 30])
    add_choice(contexts, "C", Language.values)

    guide = workbook.create_sheet(_("Нұсқаулық"))
    guide.column_dimensions["A"].width = 14
    guide.column_dimensions["B"].width = 40
    guide.column_dimensions["C"].width = 60
    lines = [
        [_("Сұрақтарды импорттау: %(subject)s") % {"subject": subject.name}],
        [],
        [_("«%(sheet)s» парағы: бір жол — бір сұрақ.") % {"sheet": QUESTIONS_SHEET}],
        [_("Тақырыпша, деңгей, тіл, мәтін, 4 нұсқа және дұрыс жауап — міндетті.")],
        [_("Контекстік сұрақта «Контекст кілті» бағанына контекстің кілтін жазыңыз (мысалы, K1).")],
        [
            _(
                "«%(sheet)s» парағы: кілт, атауы, тіл, мәтін, код, сурет. "
                "Бір контекстке ең көбі %(max)s сұрақ."
            )
            % {"sheet": CONTEXTS_SHEET, "max": QUESTIONS_PER_CONTEXT}
        ],
        [
            _(
                "Сурет: Кірістіру → Суреттер → Ұяшықтардың үстіне орналастыру. Суреттің сол жақ "
                "жоғарғы бұрышын ұяшыққа қойыңыз: сұрақ пен контекст суреті — «Сурет» бағанына, "
                "жауап суреті — сол жауаптың (A–D) ұяшығына."
            )
        ],
        [
            _("Суреті бар жауаптың мәтіні бос болуы мүмкін. Сурет — jpg, png немесе webp, ең көбі %(size)s МБ.")
            % {"size": MAX_IMAGE_MB}
        ],
        [_("Қате болса, ештеңе сақталмайды. Банкте бұрыннан бар сұрақ қайта қосылмайды.")],
        [],
        [_("Мысал:")],
        question_headers(),
        ["03", "B", "kk", "", "Алғашқы ЭЕМ қалай аталады?", "", "ENIAC", "IBM PC", "Apple I", "Altair", "A", ""],
        [],
        [_("Пәннің тақырыпшалары:")],
        [_("Тақырыпша"), _("Тақырып"), _("Атауы")],
    ]
    for line in lines:
        guide.append(line)
    for subtopic in Subtopic.objects.filter(topic__subject=subject).select_related("topic"):
        guide.append([f"{subtopic.number:02d}", str(subtopic.topic), subtopic.name])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
